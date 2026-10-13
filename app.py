from __future__ import annotations
import os, re, tempfile, subprocess, json, secrets, hashlib
from pathlib import Path
from datetime import timedelta
from fastapi import FastAPI, UploadFile, File, HTTPException, Request
from fastapi.responses import HTMLResponse, JSONResponse
from pydantic import BaseModel
import fitz
from pptx import Presentation
from docx import Document
import engine as engine_module
from engine import add_document, ask, make_quiz, record_answer, record_stored_answer, topics_summary, concept_summary, evaluation, clear_data, get_chunks

ROOT=Path(__file__).parent
app=FastAPI(title='LUMEN ALS — Adaptive Learning System', version='1.1.0')

SESSION_COOKIE = 'lumen_session'

@app.middleware('http')
async def isolate_browser_workspaces(request: Request, call_next):
    """Optional per-browser SQLite isolation for hosted deployments.

    The cookie contains a random bearer identifier, never a filesystem path. Only its
    SHA-256 digest is used as the session database filename.
    """
    if os.environ.get('LUMEN_SESSION_ISOLATION', '0') != '1':
        return await call_next(request)
    sid = request.cookies.get(SESSION_COOKIE, '')
    if not re.fullmatch(r'[A-Za-z0-9_-]{40,100}', sid):
        sid = secrets.token_urlsafe(32)
        is_new = True
    else:
        is_new = False
    data_dir = Path(os.environ.get('LUMEN_DATA_DIR', str(ROOT / 'data')))
    db_path = data_dir / 'sessions' / (hashlib.sha256(sid.encode('utf-8')).hexdigest() + '.sqlite3')
    token = engine_module.set_request_db_path(str(db_path))
    try:
        response = await call_next(request)
        if is_new:
            response.set_cookie(SESSION_COOKIE, sid, max_age=60*60*24*30, httponly=True,
                                secure=os.environ.get('LUMEN_COOKIE_SECURE', '0') == '1',
                                samesite='lax', path='/')
        response.headers['Cache-Control'] = 'no-store'
        return response
    finally:
        engine_module.reset_request_db_path(token)

@app.get('/', response_class=HTMLResponse)
def home(): return (ROOT/'web'/'index.html').read_text(encoding='utf-8')

@app.get('/api/state')
def state():
    chunks=get_chunks()
    sources=[]
    for c in chunks:
        if c['source'] not in sources: sources.append(c['source'])
    return {'sources':sources,'chunks':len(chunks),'topics':topics_summary(),'concept_map':concept_summary(),'evaluation':evaluation()}

@app.post('/api/upload')
async def upload(files: list[UploadFile] = File(...)):
    if not files:
        raise HTTPException(400, 'Choose at least one file.')
    if len(files) > 10:
        raise HTTPException(413, 'Upload at most 10 files per request.')
    result=[]
    total_bytes = 0
    for f in files:
        name=Path(f.filename or 'upload').name
        suffix=Path(name).suffix.lower()
        raw=await f.read()
        total_bytes += len(raw)
        if len(raw)>25*1024*1024:
            result.append({'name':name,'ok':False,'message':'File exceeds the 25 MB per-file limit.'}); continue
        if total_bytes>50*1024*1024:
            result.append({'name':name,'ok':False,'message':'Request exceeds the 50 MB total upload limit.'}); break
        pages=[]; kind='document'
        try:
            if suffix=='.pdf':
                doc=fitz.open(stream=raw,filetype='pdf')
                for i,page in enumerate(doc):
                    txt=page.get_text('text')
                    txt='\n'.join(line for line in txt.splitlines() if not re.search(r'LUMEN sample course|^Page \d+$', line.strip(), re.I))
                    # OCR embedded figures when the local Tesseract executable is available.
                    try:
                        import pytesseract
                        from PIL import Image
                        for img in page.get_images(full=True):
                            pix=fitz.Pixmap(doc,img[0])
                            if pix.width*pix.height>16000000: continue
                            if pix.n>4: pix=fitz.Pixmap(fitz.csRGB,pix)
                            ocr=pytesseract.image_to_string(Image.frombytes('RGB' if pix.n>=3 else 'L',[pix.width,pix.height],pix.samples))
                            if len(ocr.strip())>20: txt += '\nFigure OCR: '+ocr
                    except Exception: pass
                    pages.append((f'Page {i+1}',txt))
                kind='pdf'
            elif suffix=='.pptx':
                prs=Presentation(__import__('io').BytesIO(raw))
                for i,slide in enumerate(prs.slides):
                    parts=[]
                    for shape in slide.shapes:
                        if getattr(shape,'has_text_frame',False) and shape.text.strip(): parts.append(shape.text)
                        if getattr(shape,'has_table',False):
                            parts.extend(' | '.join(c.text for c in row.cells) for row in shape.table.rows)
                        # OCR text embedded in slide images when the optional local OCR stack is available.
                        if getattr(shape,'shape_type',None)==13:
                            try:
                                import pytesseract
                                from PIL import Image
                                image=Image.open(__import__('io').BytesIO(shape.image.blob)).convert('RGB')
                                ocr=pytesseract.image_to_string(image)
                                if len(ocr.strip())>15: parts.append('Figure OCR: '+ocr)
                            except Exception: pass
                    pages.append((f'Slide {i+1}','\n'.join(parts)))
                kind='slides'
            elif suffix=='.docx':
                doc=Document(__import__('io').BytesIO(raw))
                pages=[(f'Paragraph {i+1}',p.text) for i,p in enumerate(doc.paragraphs) if p.text.strip()]
                for ti,t in enumerate(doc.tables):
                    pages.append((f'Table {ti+1}','\n'.join(' | '.join(c.text for c in row.cells) for row in t.rows)))
                kind='document'
            elif suffix in ('.txt','.md'):
                text=raw.decode('utf-8-sig',errors='replace'); pages=[(f'Line {i+1}',line) for i,line in enumerate(text.splitlines()) if line.strip()]; kind='text'
            elif suffix in ('.srt','.vtt'):
                text=raw.decode('utf-8-sig',errors='replace'); blocks=re.split(r'\n\s*\n',text); kind='transcript'
                for b in blocks:
                    lines=[x.strip() for x in b.splitlines() if x.strip()]
                    time=next((x for x in lines if '-->' in x),None)
                    content=[x for x in lines if '-->' not in x and not re.fullmatch(r'\d+',x) and x!='WEBVTT']
                    if content: pages.append((f'Timestamp {time.split(" --> ")[0]}' if time else 'Transcript segment',' '.join(content)))
            elif suffix in ('.mp4','.mov','.mkv','.webm','.mp3','.wav','.m4a'):
                # Transcription is optional to keep the base install small and free.
                try:
                    from faster_whisper import WhisperModel
                    with tempfile.TemporaryDirectory() as td:
                        path=Path(td)/name; path.write_bytes(raw)
                        model=WhisperModel(os.environ.get('LUMEN_WHISPER_MODEL','base'),device='cpu',compute_type='int8')
                        segments,_=model.transcribe(str(path),vad_filter=True)
                        pages=[(f'Timestamp {timedelta(seconds=int(s.start))}',s.text) for s in segments]
                    kind='video transcript'
                except ImportError:
                    result.append({'name':name,'ok':False,'message':'Video/audio transcription needs the optional free local package faster-whisper. Install it with: pip install faster-whisper. Alternatively upload a .vtt or .srt transcript for timestamped indexing.'}); continue
                except Exception as e:
                    result.append({'name':name,'ok':False,'message':f'Transcription failed: {str(e)[:180]}. Try a .vtt/.srt transcript instead.'}); continue
            else:
                result.append({'name':name,'ok':False,'message':'Unsupported format. Use PDF, PPTX, DOCX, TXT, MD, SRT, VTT, or audio/video with optional transcription.'}); continue
            count=add_document(name,kind,pages)
            result.append({'name':name,'ok':count>0,'chunks_added':count,'pages_or_segments':len(pages),'message':f'Indexed {count} source chunks.' if count else 'No new usable text found. The file may be image-only, empty, or already indexed.'})
        except Exception as e:
            result.append({'name':name,'ok':False,'message':f'Could not parse file: {str(e)[:180]}'})
    return {'results':result,'state':state()}

class AskRequest(BaseModel): question:str
@app.post('/api/ask')
def api_ask(req:AskRequest):
    if len(req.question.strip())<3: raise HTTPException(400,'Ask a more specific question.')
    return ask(req.question[:1000])

class QuizRequest(BaseModel): count:int=5
@app.post('/api/quiz')
def api_quiz(req:QuizRequest):
    if not get_chunks(): raise HTTPException(400,'Upload course material first.')
    qs=make_quiz(max(1,min(req.count,10)))
    # Never send answer keys, grading keywords, or explanations to the browser before submission.
    public_keys={'id','type','question','options','topic','source','locator','mastery_before','selection_reason'}
    public_questions=[{k:v for k,v in q.items() if k in public_keys} for q in qs]
    return {'questions':public_questions,'message':None if qs else 'Not enough source sentences to generate safe cloze questions. Upload fuller notes or slides.'}

class AnswerRequest(BaseModel): question_id:str; answer:str
@app.post('/api/answer')
def api_answer(req:AnswerRequest):
    if not req.question_id.strip(): raise HTTPException(400,'Missing question identifier.')
    if len(req.answer) > 2000: raise HTTPException(413, 'Answer is too long (maximum 2,000 characters).')
    try: return record_stored_answer(req.question_id,req.answer)
    except ValueError as e: raise HTTPException(404,str(e))

@app.get('/api/evaluation')
def api_evaluation(): return evaluation()

@app.post('/api/reset')
def reset():
    clear_data(); return {'ok':True,'state':state()}
