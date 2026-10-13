from pathlib import Path
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
ROOT=Path(__file__).parent/'sample_materials'; ROOT.mkdir(exist_ok=True)
# Author-created sample material; not presented as external course content.
sections=[
('Newton’s Laws: A Short Learning Pack','This original sample course pack is included to demonstrate source ingestion, citations, tutoring, and assessment. It is not copied from an external textbook.'),
('1. Newton’s First Law','An object remains at rest, or continues moving at constant velocity in a straight line, unless a net external force acts on it. This tendency to resist a change in motion is called inertia. Mass is a measure of an object’s inertia. A book resting on a table stays at rest because the forces on it are balanced: gravity pulls downward and the table’s normal force pushes upward.'),
('2. Newton’s Second Law','Newton’s second law states that the net force on an object equals its mass multiplied by its acceleration. The relationship is commonly written as F = m × a. If the net force increases while mass remains constant, acceleration increases. If mass increases while net force remains constant, acceleration decreases. Force is measured in newtons, mass in kilograms, and acceleration in metres per second squared.'),
('3. Newton’s Third Law','When one object exerts a force on a second object, the second object exerts a force of equal magnitude and opposite direction on the first. These action–reaction forces act on different objects, so they do not cancel each other on a single object. A swimmer pushes water backward, and the water pushes the swimmer forward.'),
('4. Balanced and Unbalanced Forces','Forces are balanced when their vector sum is zero. An object then has no acceleration: it may remain at rest or continue at constant velocity. Forces are unbalanced when the net force is non-zero, causing acceleration in the direction of the net force. The net force is the vector sum of all forces, not simply the largest force.'),
('5. Worked Example','A 2 kg cart experiences a net horizontal force of 6 N. Using F = m × a, acceleration equals F divided by m: a = 6 N / 2 kg = 3 m/s². If the same 6 N net force acts on a 4 kg cart, its acceleration is 1.5 m/s². Before applying the second law, first understand net force. Always identify the net force before applying the second law.'),
('6. Common Misconceptions','A force is not required to keep an object moving at constant velocity; a net force is required to change velocity. Action–reaction forces do not cancel because they act on different objects. Mass and weight are not identical: mass measures inertia, while weight is the gravitational force acting on an object. A diagram should label forces acting on the chosen object and show their directions.')]
styles=getSampleStyleSheet(); styles.add(ParagraphStyle(name='TitleCenter',parent=styles['Title'],alignment=TA_CENTER,textColor=colors.HexColor('#9c7c39'),fontSize=22,leading=28,spaceAfter=14)); styles.add(ParagraphStyle(name='Section',parent=styles['Heading2'],textColor=colors.HexColor('#8d6e2f'),fontSize=14,leading=18,spaceBefore=14,spaceAfter=7)); styles.add(ParagraphStyle(name='BodyX',parent=styles['BodyText'],fontSize=10.5,leading=16,spaceAfter=8))
story=[Paragraph(sections[0][0],styles['TitleCenter']),Paragraph(sections[0][1],styles['BodyX']),Spacer(1,14)]
for h,t in sections[1:]: story.extend([Paragraph(h,styles['Section']),Paragraph(t,styles['BodyX'])])
def footer(canvas,doc):
 canvas.saveState(); canvas.setStrokeColor(colors.HexColor('#d8b56d')); canvas.line(42,38,570,38); canvas.setFont('Helvetica',8); canvas.setFillColor(colors.HexColor('#6c747a')); canvas.drawString(42,24,'LUMEN sample course · Original demo material'); canvas.drawRightString(570,24,f'Page {doc.page}'); canvas.restoreState()
SimpleDocTemplate(str(ROOT/'newtons_laws_sample.pdf'),pagesize=letter,rightMargin=48,leftMargin=48,topMargin=48,bottomMargin=52,title='Newton’s Laws — LUMEN Sample Course').build(story,onFirstPage=footer,onLaterPages=footer)
prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
def slide(title, bullets, subtitle=None):
 s=prs.slides.add_slide(prs.slide_layouts[6]); bg=s.background.fill; bg.solid(); bg.fore_color.rgb=RGBColor(12,16,20)
 box=s.shapes.add_textbox(Inches(.8), Inches(.55), Inches(11.7), Inches(1.0)); tf=box.text_frame; tf.text=title
 p=tf.paragraphs[0]; p.font.name='Aptos Display'; p.font.size=Pt(30); p.font.bold=True; p.font.color.rgb=RGBColor(216,181,109)
 if subtitle:
  b=s.shapes.add_textbox(Inches(.85), Inches(1.55), Inches(11.5), Inches(.55)); p=b.text_frame.paragraphs[0]; p.text=subtitle; p.font.size=Pt(15); p.font.color.rgb=RGBColor(151,166,177)
 b=s.shapes.add_textbox(Inches(1.0), Inches(2.3 if subtitle else 1.8), Inches(11.2), Inches(4.5)); tf=b.text_frame; tf.word_wrap=True
 for i,txt in enumerate(bullets):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text='•  '+txt; p.font.name='Aptos'; p.font.size=Pt(22); p.font.color.rgb=RGBColor(230,236,239); p.space_after=Pt(18)
slide('Newton’s Laws of Motion',['First law: inertia and constant velocity','Second law: net force determines acceleration','Third law: equal and opposite forces on different objects'],'LUMEN sample lesson · original demonstration material')
slide('Second Law: F = m × a',['Net force = mass × acceleration','6 N acting on a 2 kg cart gives 3 m/s²','Double the mass at the same net force → half the acceleration'],'Use the net force, not an individual force.')
slide('Misconceptions to test',['Motion does not require a net force','Action–reaction forces act on different objects','Mass is not the same as weight'],'Explain each answer using evidence from the source.')
# A simple source-authored visual figure with labels that the slide parser can index.
s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=RGBColor(12,16,20)
t=s.shapes.add_textbox(Inches(.8), Inches(.55), Inches(11.7), Inches(.8)); p=t.text_frame.paragraphs[0]; p.text='Force diagram: a cart on a level surface'; p.font.name='Aptos Display'; p.font.size=Pt(28); p.font.bold=True; p.font.color.rgb=RGBColor(216,181,109)
cart=s.shapes.add_shape(1, Inches(4.5), Inches(2.7), Inches(2.0), Inches(1.2)); cart.fill.solid(); cart.fill.fore_color.rgb=RGBColor(35,46,55); cart.line.color.rgb=RGBColor(216,181,109); cart.text='2 kg cart'; cart.text_frame.paragraphs[0].font.size=Pt(18); cart.text_frame.paragraphs[0].font.color.rgb=RGBColor(237,241,243)
arrow=s.shapes.add_shape(33, Inches(6.5), Inches(3.0), Inches(2.5), Inches(.5)); arrow.fill.solid(); arrow.fill.fore_color.rgb=RGBColor(216,181,109); arrow.line.color.rgb=RGBColor(216,181,109)
label=s.shapes.add_textbox(Inches(7.0), Inches(2.35), Inches(3.0), Inches(.5)); p=label.text_frame.paragraphs[0]; p.text='Net force: 6 N'; p.font.size=Pt(19); p.font.color.rgb=RGBColor(237,241,243)
result=s.shapes.add_textbox(Inches(4.5), Inches(4.4), Inches(5.5), Inches(.8)); p=result.text_frame.paragraphs[0]; p.text='Calculated acceleration: 3 m/s²'; p.font.size=Pt(22); p.font.bold=True; p.font.color.rgb=RGBColor(124,216,189)
prs.save(ROOT/'newtons_laws_sample.pptx')
(ROOT/'newtons_laws_sample.vtt').write_text('''WEBVTT\n\n00:00:00.000 --> 00:00:08.000\nNewton’s first law says an object stays at rest or continues at constant velocity unless a net external force acts on it.\n\n00:00:08.000 --> 00:00:18.000\nNewton’s second law states that net force equals mass multiplied by acceleration. We write this as F equals m times a.\n\n00:00:18.000 --> 00:00:28.000\nA six newton net force on a two kilogram cart produces an acceleration of three metres per second squared.\n\n00:00:28.000 --> 00:00:38.000\nAction and reaction forces have equal magnitude and opposite directions, but they act on different objects.\n''',encoding='utf-8')
(ROOT/'SOURCE_NOTE.txt').write_text('These are original demonstration learning materials created for the LUMEN hackathon submission. They are not claimed to be an external textbook, lecture, or public dataset. The PDF, slide deck, and transcript are designed to exercise page, slide, and timestamp citation paths.',encoding='utf-8')
print('created',*[p.name for p in ROOT.iterdir()])
